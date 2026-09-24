"""Memory rules, recall, document parsing, chunking and retrieval."""

from __future__ import annotations

import io

import pytest

from app.services import files as file_service
from app.services import memory_service as ms
from app.services.documents.chunking import chunk_document
from app.services.documents.parsing import PageText, ParsedDocument, parse_docx_bytes, parse_pdf_bytes
from app.services.documents.retrieval import rank_chunks, section_refs

# --------------------------------------------------------------------------- memory rules


@pytest.mark.parametrize("value,action", [
    ("LinkGuard AI", "store_active"),
    ("sk-ant-api03-abcdefghijklmnopqrstu", "reject"),
    ("4111 1111 1111 1111", "reject"),
    ("my password is hunter2", "reject"),
    ("1234 5678 9012", "reject"),  # Aadhaar-like
    ("I was diagnosed with anxiety", "store_pending"),
])
def test_memory_policy(value: str, action: str) -> None:
    decision = ms.evaluate(ms.MemoryCandidate("fact", "something", value))
    assert decision.action == action


def test_inferred_memories_need_approval() -> None:
    cand = ms.MemoryCandidate("fact", "favourite language", "Rust", explicit=False)
    assert ms.evaluate(cand, {"memory": {"require_approval_for_inferred": True}}).action == "store_pending"
    assert ms.evaluate(cand, {"memory": {"require_approval_for_inferred": False}}).action == "store_active"


def test_transient_states_not_stored() -> None:
    assert ms.evaluate(ms.MemoryCandidate("fact", "mood", "I'm tired")).action == "skip"


async def test_recall_resolves_project_reference(user) -> None:
    cand = ms.extract_explicit("My college AI project is LinkGuard AI at /home/me/linkguard")
    assert cand and cand.subject == "college AI project" and cand.value == "LinkGuard AI"
    await ms.save_candidate(user.id, cand)
    await ms.save_candidate(user.id, ms.MemoryCandidate("person", "sister", "Ana"))
    hits = await ms.recall(user.id, "Open my college AI project")
    assert hits and hits[0][0].value == "LinkGuard AI"
    assert hits[0][0].attributes["path"] == "/home/me/linkguard"
    assert all(m.subject != "sister" for m, _ in hits)


async def test_restating_updates_instead_of_duplicating(user) -> None:
    await ms.save_candidate(user.id, ms.MemoryCandidate("project", "college AI project", "LinkGuard"))
    await ms.save_candidate(user.id, ms.MemoryCandidate("project", "College AI Project", "LinkGuard AI"))
    rows = await ms.list_memories(user.id)
    assert len(rows) == 1 and rows[0].value == "LinkGuard AI"


async def test_update_rejects_secrets(user) -> None:
    from app.core.errors import ValidationFailedError

    _, row = await ms.save_candidate(user.id, ms.MemoryCandidate("fact", "wifi", "home network"))
    with pytest.raises(ValidationFailedError):
        await ms.update_memory(user.id, row.id, value="password: hunter2")


# --------------------------------------------------------------------------- documents


def make_pdf(pages: list[str]) -> bytes:
    """Build a small text PDF by hand (no external generator needed)."""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + len(pages) * 2
    for i, text in enumerate(pages):
        lines = text.split("\n")
        ops = ["BT", "/F1 12 Tf", "14 TL", "50 750 Td"]
        for line in lines:
            safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({safe}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode()
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + i * 2} 0 R "
                       f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


SYLLABUS = [
    "Networks Course\nUnit 1: Basics\nA network connects computers together.",
    "Unit 2: Protocols\nTCP provides reliable delivery with a three-way handshake.",
    "UNIT III: Routing\nRouting chooses paths. Dijkstra computes shortest paths.\nOSPF uses link-state routing.",
    "Unit 4 - Security\nFirewalls filter traffic.",
]


def test_pdf_parsing_and_sections() -> None:
    doc = parse_pdf_bytes(make_pdf(SYLLABUS))
    assert doc.meta["page_count"] == 4 and "three-way handshake" in doc.text
    chunks, sections = chunk_document(doc)
    titles = [s["title"] for s in sections]
    assert any("Unit 2" in t for t in titles) and any("UNIT III" in t for t in titles)
    assert all(c.page for c in chunks)


def test_unit_reference_retrieves_that_unit_only() -> None:
    doc = parse_pdf_bytes(make_pdf(SYLLABUS))
    chunks, _ = chunk_document(doc)
    results = rank_chunks(chunks, "Explain Unit 3 simply")
    assert results and all("Routing" in r.chunk.text or "Dijkstra" in r.chunk.text for r in results)
    assert section_refs("explain unit iii and chapter 2") == {("unit", 3), ("chapter", 2)}


def test_bm25_ranking_without_section() -> None:
    doc = ParsedDocument(pages=[PageText(1, "Photosynthesis happens in chloroplasts.\n\n" + "Filler text. " * 200 +
                                         "\n\nMitochondria produce ATP through respiration.")])
    chunks, _ = chunk_document(doc, target_chars=300)
    top = rank_chunks(chunks, "where is ATP produced?", top_k=1)
    assert "ATP" in top[0].chunk.text


def test_docx_headings_become_sections() -> None:
    import docx

    d = docx.Document()
    d.add_heading("Chapter 1 Intro", 1)
    d.add_paragraph("Intro text")
    d.add_heading("Chapter 2 Methods", 1)
    d.add_paragraph("We used surveys.")
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_docx_bytes(buf.getvalue())
    chunks, sections = chunk_document(doc)
    assert [s["title"] for s in sections] == ["Chapter 1 Intro", "Chapter 2 Methods"]
    assert rank_chunks(chunks, "chapter 2")[0].chunk.text.endswith("We used surveys.")


def test_scanned_pdf_is_explained() -> None:
    from app.core.errors import ValidationFailedError

    with pytest.raises(ValidationFailedError) as exc:
        parse_pdf_bytes(make_pdf(["", ""]))
    assert exc.value.info.code == "pdf_no_text"


async def test_upload_pipeline(user) -> None:
    row = await file_service.store_upload(user.id, "../../syllabus.pdf", make_pdf(SYLLABUS))
    assert row.status == "ready" and row.filename == "syllabus.pdf" and row.page_count == 4
    assert ".." not in row.stored_name and row.stored_name.endswith(".pdf")
    results = await file_service.search(user.id, "unit 2 handshake")
    assert results and "handshake" in results[0].chunk.text
    await file_service.delete_file(user.id, row.id)
    assert await file_service.list_files(user.id) == []
