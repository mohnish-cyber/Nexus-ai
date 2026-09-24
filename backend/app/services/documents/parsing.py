"""Text extraction for PDF, DOCX, text and code files."""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field

from app.core.errors import ValidationFailedError

logger = logging.getLogger(__name__)

MAX_PDF_PAGES = 1500
MAX_TEXT_CHARS = 5_000_000

# "Unit 3", "UNIT III: Networks", "Chapter 4 - Graphs", "Module-2", "Question 5", "Q5."
HEADING_RE = re.compile(
    r"^\s*(?P<kind>unit|chapter|module|section|part|lesson|lecture|week|topic|question|q)"
    r"\s*[-:.#]?\s*(?P<num>\d{1,3}|[ivxlc]{1,7})\b[\s:.\-–—)]*(?P<title>.{0,120})$",
    re.IGNORECASE,
)
MD_HEADING_RE = re.compile(r"^\s{0,3}(#{1,4})\s+(?P<title>.{1,150})$")
CODE_SYMBOL_RE = re.compile(
    r"^\s*(?:export\s+)?(?:default\s+)?(?:async\s+)?(?:public\s+|private\s+|protected\s+|static\s+)*"
    r"(?P<kw>class|def|function|func|fn|interface|struct|enum)\s+(?P<name>[A-Za-z_][A-Za-z0-9_]*)"
)


@dataclass
class PageText:
    number: int
    text: str


@dataclass
class ParsedDocument:
    pages: list[PageText]
    title: str | None = None
    kind: str = "text"
    meta: dict = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)


def detect_heading(line: str, kind: str = "text") -> str | None:
    stripped = line.strip()
    if not stripped or len(stripped) > 160:
        return None
    if kind == "code":
        m = CODE_SYMBOL_RE.match(line)
        return f"{m.group('kw')} {m.group('name')}" if m else None
    m = MD_HEADING_RE.match(line)
    if m:
        return m.group("title").strip("# ").strip()
    m = HEADING_RE.match(stripped)
    if m:
        # Avoid matching sentences like "Part of the reason is..."
        title = m.group("title").strip()
        if m.group("kind").lower() in ("part", "section", "topic") and title and title[:1].islower():
            return None
        return stripped
    return None


def parse_pdf_bytes(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ValidationFailedError(
                    "This PDF is password-protected.", code="pdf_encrypted",
                    next_step="Remove the password and upload it again.") from exc
        pages: list[PageText] = []
        for i, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:  # malformed page: skip, keep going
                logger.info("PDF page %s could not be read: %s", i, exc)
                text = ""
            pages.append(PageText(i, text))
        title = None
        try:
            title = (reader.metadata.title if reader.metadata else None) or None
        except Exception:
            title = None
    except ValidationFailedError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise ValidationFailedError("This PDF could not be read; it may be corrupted.", code="pdf_unreadable",
                                    reason=str(exc)[:200]) from exc
    total = sum(len(p.text.strip()) for p in pages)
    if pages and total < 20:
        raise ValidationFailedError(
            "No text could be extracted from this PDF.",
            code="pdf_no_text",
            reason="It is probably a scanned document (images only). OCR is not enabled.",
            next_step="Upload a text-based PDF, or upload page images so NEXUS can read them with vision.",
        )
    return ParsedDocument(pages=pages, title=title, kind="pdf", meta={"page_count": len(reader.pages)})


def parse_docx_bytes(data: bytes) -> ParsedDocument:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ValidationFailedError("This Word document could not be read.", code="docx_unreadable") from exc
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style == "title":
            level = re.sub(r"\D", "", style) or "1"
            lines.append(f"{'#' * min(int(level), 4)} {text}")
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    title = document.core_properties.title or None
    return ParsedDocument(pages=[PageText(1, "\n".join(lines))], title=title, kind="docx")


def parse_text_bytes(data: bytes, kind: str = "text") -> ParsedDocument:
    text = data[:MAX_TEXT_CHARS].decode("utf-8", errors="replace")
    return ParsedDocument(pages=[PageText(1, text)], kind=kind)


def parse_document(kind: str, data: bytes) -> ParsedDocument:
    if kind == "pdf":
        return parse_pdf_bytes(data)
    if kind == "docx":
        return parse_docx_bytes(data)
    if kind in ("text", "code"):
        return parse_text_bytes(data, kind)
    raise ValidationFailedError(f"Cannot extract text from {kind} files.", code="unsupported_file_type")
